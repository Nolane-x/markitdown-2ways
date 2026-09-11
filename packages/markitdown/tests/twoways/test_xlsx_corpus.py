from __future__ import annotations

from io import BytesIO
from pathlib import Path

import openpyxl

from markitdown.twoways.capabilities import build_capability_report
from markitdown.twoways.formats.xlsx.reader import read_xlsx_ir
from markitdown.twoways.formats.xlsx.writer import patch_xlsx
from markitdown.twoways.ir.edits import EditOperation

from ._xlsx_fixtures import SHEET1, make_xlsx


TEST_FILES = Path(__file__).resolve().parents[1] / "test_files"


def test_patched_workbook_opens_with_openpyxl_and_preserves_untouched_semantics() -> (
    None
):
    sheet1 = SHEET1.replace(' t="s" s="1"', ' t="s"')
    source = make_xlsx(replacements={"xl/worksheets/sheet1.xml": sheet1})
    document = read_xlsx_ir(BytesIO(source))
    sheet_id = document.canvases[0].root_node_ids[0]
    edit = EditOperation(
        operation_id="differential-open-edit",
        type="update_sheet_cells",
        target_node_id=sheet_id,
        payload={
            "cells": [
                {"row": 0, "column": 1, "old_value": 7, "value": 13.5},
                {"row": 1, "column": 0, "old_value": True, "value": False},
            ]
        },
    )
    output = BytesIO()

    patch_xlsx(document, BytesIO(source), output, edits=(edit,))
    output_bytes = output.getvalue()

    workbook = openpyxl.load_workbook(
        BytesIO(output_bytes),
        read_only=True,
        data_only=False,
    )
    try:
        assert workbook.sheetnames == ["Data", "Other"]
        assert workbook["Data"]["A1"].value == "North"
        assert workbook["Data"]["B1"].value == 13.5
        assert workbook["Data"]["A2"].value is False
        assert workbook["Data"]["C1"].value == "=B1+1"
        assert workbook["Other"]["A1"].value == "untouched"
    finally:
        workbook.close()

    cached = openpyxl.load_workbook(
        BytesIO(output_bytes),
        read_only=True,
        data_only=True,
    )
    try:
        assert cached["Data"]["C1"].value == 8
    finally:
        cached.close()


def test_real_xlsx_fixture_reads_reports_capabilities_and_noops_byte_identically() -> (
    None
):
    source = (TEST_FILES / "test.xlsx").read_bytes()
    document = read_xlsx_ir(BytesIO(source))
    report = build_capability_report(document)

    assert document.canvases
    assert all(canvas.kind == "worksheet" for canvas in document.canvases)
    assert report.total_nodes == len(document.nodes)
    assert report.total_nodes > 0
    assert (
        report.writable_nodes + report.read_only_nodes + report.derived_nodes
        == report.total_nodes
    )

    output = BytesIO()
    patch_xlsx(document, BytesIO(source), output, edits=())
    assert output.getvalue() == source
