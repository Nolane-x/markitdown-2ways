from __future__ import annotations

from io import BytesIO
from zipfile import ZipFile

import pytest
from pptx import Presentation

from markitdown.twoways import PatchPreconditionError, UnsupportedEditError
from markitdown.twoways.ir.edits import EditOperation, EditPrecondition
from markitdown.twoways.ir.nodes import TablePayload
from markitdown.twoways.ir.semantics import (
    native_locator_digest,
    node_semantic_digest,
    node_semantic_text,
)

from ._pptx_fixtures import make_pptx_bytes


def _table(document):
    return next(node for node in document.nodes.values() if node.kind == "table")


def _edit(table, cells):
    return EditOperation(
        operation_id="pptx-update-table-cells",
        type="update_table_cells",
        target_node_id=table.node_id,
        precondition=EditPrecondition(
            expected_semantic_digest=node_semantic_digest(table),
            expected_native_locator_digest=native_locator_digest(table),
            expected_old_value=node_semantic_text(table),
        ),
        payload={"cells": cells},
    )


def _assert_only_member_changed(before: bytes, after: bytes, changed: str):
    with ZipFile(BytesIO(before)) as a, ZipFile(BytesIO(after)) as b:
        assert a.namelist() == b.namelist()
        for name in a.namelist():
            if name == changed:
                assert a.read(name) != b.read(name)
            else:
                assert a.read(name) == b.read(name), name


def _cell_text(payload: TablePayload, row: int, column: int) -> str:
    return next(
        (cell.text or "")
        for cell in payload.cells
        if cell.row == row and cell.column == column
    )


def test_pptx_reader_marks_fixture_table_cell_patchable():
    from markitdown.twoways.formats.pptx import read_pptx_ir

    document = read_pptx_ir(BytesIO(make_pptx_bytes()))
    table = _table(document)

    assert table.metadata["pptx:patch_capabilities"] == ("update_table_cells",)


def test_pptx_single_cell_patch_changes_only_owning_slide_and_reads_back():
    from markitdown.twoways.formats.pptx import patch_pptx, read_pptx_ir

    source = make_pptx_bytes()
    document = read_pptx_ir(BytesIO(source))
    table = _table(document)
    output = BytesIO()

    result = patch_pptx(
        document,
        BytesIO(source),
        output,
        edits=(
            _edit(
                table,
                [{"row": 1, "column": 1, "old_text": "42", "text": "47"}],
            ),
        ),
    )

    patched = output.getvalue()
    _assert_only_member_changed(source, patched, "ppt/slides/slide2.xml")
    assert result.metadata["touched_parts"] == ("ppt/slides/slide2.xml",)
    assert result.fidelity.claimed_tier == "high"

    presentation = Presentation(BytesIO(patched))
    shape = next(item for item in presentation.slides[1].shapes if item.has_table)
    assert shape.table.cell(1, 1).text == "47"

    reread = read_pptx_ir(BytesIO(patched))
    payload = reread.nodes[table.node_id].payload
    assert isinstance(payload, TablePayload)
    assert _cell_text(payload, 1, 1) == "47"
    assert _cell_text(payload, 1, 0) == "APAC"


def test_pptx_multi_cell_patch_preserves_unedited_cells():
    from markitdown.twoways.formats.pptx import patch_pptx, read_pptx_ir

    source = make_pptx_bytes()
    document = read_pptx_ir(BytesIO(source))
    table = _table(document)
    output = BytesIO()
    patch_pptx(
        document,
        BytesIO(source),
        output,
        edits=(
            _edit(
                table,
                [
                    {"row": 0, "column": 0, "old_text": "Region", "text": "Market"},
                    {"row": 1, "column": 1, "old_text": "42", "text": "47"},
                ],
            ),
        ),
    )
    reread = read_pptx_ir(BytesIO(output.getvalue()))
    payload = reread.nodes[table.node_id].payload
    assert isinstance(payload, TablePayload)
    assert _cell_text(payload, 0, 0) == "Market"
    assert _cell_text(payload, 0, 1) == "Revenue"
    assert _cell_text(payload, 1, 0) == "APAC"
    assert _cell_text(payload, 1, 1) == "47"


def test_pptx_table_patch_rejects_stale_old_text_before_output():
    from markitdown.twoways.formats.pptx import patch_pptx, read_pptx_ir

    source = make_pptx_bytes()
    document = read_pptx_ir(BytesIO(source))
    table = _table(document)
    output = BytesIO()

    with pytest.raises(PatchPreconditionError):
        patch_pptx(
            document,
            BytesIO(source),
            output,
            edits=(
                _edit(
                    table,
                    [{"row": 1, "column": 1, "old_text": "41", "text": "47"}],
                ),
            ),
        )
    assert output.getvalue() == b""


def test_pptx_table_patch_rejects_duplicate_coordinate():
    from markitdown.twoways.formats.pptx import patch_pptx, read_pptx_ir

    source = make_pptx_bytes()
    document = read_pptx_ir(BytesIO(source))
    table = _table(document)

    with pytest.raises(UnsupportedEditError):
        patch_pptx(
            document,
            BytesIO(source),
            BytesIO(),
            edits=(
                _edit(
                    table,
                    [
                        {"row": 1, "column": 1, "old_text": "42", "text": "47"},
                        {"row": 1, "column": 1, "old_text": "42", "text": "48"},
                    ],
                ),
            ),
        )
