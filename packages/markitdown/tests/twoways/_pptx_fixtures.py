from __future__ import annotations

import base64
from io import BytesIO

from pptx import Presentation
from pptx.chart.data import ChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Inches, Pt


PNG_1X1 = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mNk+A8AAQUBAScY42YAAAAASUVORK5CYII="
)


def make_pptx_bytes(
    *,
    metric: str = "38%",
    alt_text: str = "Revenue icon",
    notes_text: str = "Speaker note 38%",
) -> bytes:
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[5])
    slide.shapes.title.text = "Quarterly Revenue"

    textbox = slide.shapes.add_textbox(Inches(1), Inches(1.5), Inches(5), Inches(1))
    paragraph = textbox.text_frame.paragraphs[0]
    run1 = paragraph.add_run()
    run1.text = "Revenue "
    run1.font.bold = True
    run1.font.size = Pt(24)
    run1.font.name = "Aptos"
    run1.font.color.rgb = RGBColor(10, 20, 30)
    run2 = paragraph.add_run()
    run2.text = metric
    run2.font.italic = True
    run2.font.size = Pt(24)

    picture = slide.shapes.add_picture(BytesIO(PNG_1X1), Inches(1), Inches(3))
    picture._element._nvXxPr.cNvPr.set("descr", alt_text)
    slide.notes_slide.notes_text_frame.text = notes_text

    slide2 = presentation.slides.add_slide(presentation.slide_layouts[5])
    slide2.shapes.title.text = "Second slide"
    shape = slide2.shapes.add_shape(
        1,
        Inches(0.5),
        Inches(1.5),
        Inches(2),
        Inches(1),
    )
    shape.text = "Read-only shape"

    table_shape = slide2.shapes.add_table(
        2,
        2,
        Inches(3),
        Inches(1.5),
        Inches(3),
        Inches(1.5),
    )
    table_shape.table.cell(0, 0).text = "Region"
    table_shape.table.cell(0, 1).text = "Revenue"
    table_shape.table.cell(1, 0).text = "APAC"
    table_shape.table.cell(1, 1).text = "42"

    chart_data = ChartData()
    chart_data.categories = ["Q1", "Q2"]
    chart_data.add_series("Revenue", (38, 42))
    slide2.shapes.add_chart(
        XL_CHART_TYPE.COLUMN_CLUSTERED,
        Inches(0.5),
        Inches(3),
        Inches(5),
        Inches(2.5),
        chart_data,
    )

    output = BytesIO()
    presentation.save(output)
    return output.getvalue()


def make_grouped_pptx_bytes(*, metric: str = "38%") -> bytes:
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])

    label = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(2), Inches(0.6))
    label.text = "Grouped Revenue"
    metric_box = slide.shapes.add_textbox(
        Inches(1), Inches(1.8), Inches(2), Inches(0.6)
    )
    paragraph = metric_box.text_frame.paragraphs[0]
    paragraph.clear()
    run = paragraph.add_run()
    run.text = metric
    run.font.bold = True

    group = slide.shapes.add_group_shape([label, metric_box])
    group.name = "Revenue Group"

    output = BytesIO()
    presentation.save(output)
    return output.getvalue()
