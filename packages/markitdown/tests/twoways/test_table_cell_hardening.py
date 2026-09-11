from __future__ import annotations

from io import BytesIO

import pytest

from markitdown.twoways import UnsupportedEditError
from markitdown.twoways.ir.edits import EditOperation, EditPrecondition
from markitdown.twoways.ir.nodes import TableCell, TablePayload
from markitdown.twoways.ir.semantics import (
    native_locator_digest,
    node_semantic_digest,
    node_semantic_text,
)
from markitdown.twoways.ir.table_edits import validate_table_cell_updates

from ._docx_fixtures import build_docx_fixture
from ._pptx_fixtures import make_pptx_bytes


def _payload() -> TablePayload:
    return TablePayload(
        rows=2,
        columns=2,
        cells=(
            TableCell(row=0, column=0, text="A"),
            TableCell(row=0, column=1, text="B"),
            TableCell(row=1, column=0, text="C"),
            TableCell(row=1, column=1, text="D"),
        ),
    )


def _table(document):
    return next(node for node in document.nodes.values() if node.kind == "table")


def _edit(table, cells):
    return EditOperation(
        operation_id="hardening-table-edit",
        type="update_table_cells",
        target_node_id=table.node_id,
        precondition=EditPrecondition(
            expected_semantic_digest=node_semantic_digest(table),
            expected_native_locator_digest=native_locator_digest(table),
            expected_old_value=node_semantic_text(table),
        ),
        payload={"cells": cells},
    )


def _mutate_docx(mutator) -> bytes:
    from docx import Document

    document = Document(BytesIO(build_docx_fixture()))
    mutator(document.tables[0])
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def _mutate_pptx(mutator) -> bytes:
    from pptx import Presentation

    presentation = Presentation(BytesIO(make_pptx_bytes()))
    shape = next(item for item in presentation.slides[1].shapes if item.has_table)
    mutator(shape.table)
    output = BytesIO()
    presentation.save(output)
    return output.getvalue()


@pytest.mark.parametrize(
    "cells",
    [
        [{"row": True, "column": 0, "old_text": "A", "text": "X"}],
        [{"row": 0, "column": False, "old_text": "A", "text": "X"}],
        [{"row": 7, "column": 0, "old_text": "A", "text": "X"}],
        [{"row": 0, "column": 0, "old_text": "A", "text": "A"}],
    ],
)
def test_shared_table_edit_validation_rejects_ambiguous_or_noop_updates(cells):
    with pytest.raises(UnsupportedEditError):
        validate_table_cell_updates(_payload(), cells, format_label="TEST")


def test_docx_nested_table_is_read_only():
    from markitdown.twoways.formats.docx import read_docx_ir

    source = _mutate_docx(lambda table: table.cell(0, 0).add_table(rows=1, cols=1))
    table = _table(read_docx_ir(BytesIO(source)))

    assert table.metadata["docx:patch_capabilities"] == ()


def test_docx_multi_paragraph_cell_is_read_only():
    from markitdown.twoways.formats.docx import read_docx_ir

    source = _mutate_docx(lambda table: table.cell(0, 0).add_paragraph("extra"))
    table = _table(read_docx_ir(BytesIO(source)))

    assert table.metadata["docx:patch_capabilities"] == ()


def test_docx_merged_cell_is_read_only():
    from markitdown.twoways.formats.docx import read_docx_ir

    source = _mutate_docx(lambda table: table.cell(0, 0).merge(table.cell(0, 1)))
    table = _table(read_docx_ir(BytesIO(source)))

    assert table.metadata["docx:patch_capabilities"] == ()


def test_pptx_multi_paragraph_cell_is_read_only():
    from markitdown.twoways.formats.pptx import read_pptx_ir

    def mutate(table):
        table.cell(0, 0).text_frame.add_paragraph().text = "extra"

    table = _table(read_pptx_ir(BytesIO(_mutate_pptx(mutate))))

    assert table.metadata["pptx:patch_capabilities"] == ()


def test_pptx_merged_cell_is_read_only():
    from markitdown.twoways.formats.pptx import read_pptx_ir

    source = _mutate_pptx(lambda table: table.cell(0, 0).merge(table.cell(0, 1)))
    table = _table(read_pptx_ir(BytesIO(source)))

    assert table.metadata["pptx:patch_capabilities"] == ()


@pytest.mark.parametrize("new_text", ["Doanh thu 日本語 🚀", ""])
def test_docx_unicode_and_empty_cell_text_round_trip(new_text):
    from markitdown.twoways.formats.docx import patch_docx, read_docx_ir

    source = build_docx_fixture()
    document = read_docx_ir(BytesIO(source))
    table = _table(document)
    output = BytesIO()

    patch_docx(
        document,
        BytesIO(source),
        output,
        edits=(
            _edit(
                table,
                [{"row": 1, "column": 1, "old_text": "38%", "text": new_text}],
            ),
        ),
    )

    reread = read_docx_ir(BytesIO(output.getvalue()))
    patched = _table(reread)
    value = next(
        cell.text
        for cell in patched.payload.cells
        if cell.row == 1 and cell.column == 1
    )
    assert value == new_text


@pytest.mark.parametrize("new_text", ["Doanh thu 日本語 🚀", ""])
def test_pptx_unicode_and_empty_cell_text_round_trip(new_text):
    from markitdown.twoways.formats.pptx import patch_pptx, read_pptx_ir

    source = make_pptx_bytes()
    document = read_pptx_ir(BytesIO(source))
    table = _table(document)
    output = BytesIO()

    patch_pptx(
        document,
        BytesIO(source),
        output,
        edits=(
            _edit(
                table,
                [{"row": 1, "column": 1, "old_text": "42", "text": new_text}],
            ),
        ),
    )

    reread = read_pptx_ir(BytesIO(output.getvalue()))
    patched = _table(reread)
    value = next(
        cell.text
        for cell in patched.payload.cells
        if cell.row == 1 and cell.column == 1
    )
    assert value == new_text
