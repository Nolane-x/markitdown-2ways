from __future__ import annotations

from base64 import b64decode
from io import BytesIO


_PNG_1X1 = b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAusB9Wl2n1cAAAAASUVORK5CYII="
)


def _add_hyperlink(paragraph, text: str, url: str) -> str:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.opc.constants import RELATIONSHIP_TYPE as RT

    relationship_id = paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relationship_id)
    run = OxmlElement("w:r")
    run_properties = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    run_properties.extend((color, underline))
    text_node = OxmlElement("w:t")
    text_node.text = text
    run.append(run_properties)
    run.append(text_node)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)
    return relationship_id


def build_docx_fixture() -> bytes:
    from docx import Document
    from docx.shared import Inches

    document = Document()
    document.core_properties.title = "MarkItDown 2Ways DOCX Fixture"
    document.core_properties.author = "Nolane"

    paragraph = document.add_paragraph()
    first = paragraph.add_run("Revenue ")
    first.bold = True
    second = paragraph.add_run("38%")
    second.italic = True

    hyperlink_paragraph = document.add_paragraph()
    hyperlink_paragraph.add_run("Visit ")
    _add_hyperlink(hyperlink_paragraph, "OpenAI", "https://openai.com/")
    hyperlink_paragraph.add_run(" today")

    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Metric"
    table.cell(0, 1).text = "Value"
    table.cell(1, 0).text = "Revenue"
    table.cell(1, 1).text = "38%"

    picture_run = document.add_paragraph().add_run()
    inline_shape = picture_run.add_picture(BytesIO(_PNG_1X1), width=Inches(0.25))
    inline_shape._inline.docPr.set("descr", "Green status pixel")

    section = document.sections[0]
    section.header.paragraphs[0].text = "Confidential Header"
    section.footer.paragraphs[0].text = "Page Footer"

    output = BytesIO()
    document.save(output)
    return output.getvalue()
