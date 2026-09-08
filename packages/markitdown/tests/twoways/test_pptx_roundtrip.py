from __future__ import annotations

from io import BytesIO
from zipfile import ZipFile

import pytest
from pptx import Presentation

from markitdown.twoways import (
    EditOperation,
    EditPrecondition,
    ImagePayload,
    SourcePackageMismatchError,
    TextPayload,
)
from markitdown.twoways.ir.semantics import (
    native_locator_digest,
    node_semantic_digest,
    node_semantic_text,
)

from ._pptx_fixtures import make_pptx_bytes


def _text_node(document, text="Revenue 38%"):
    for node in document.nodes.values():
        if isinstance(node.payload, TextPayload) and node.payload.text == text:
            return node
    raise AssertionError("target text node not found")


def _image_node(document):
    for node in document.nodes.values():
        if isinstance(node.payload, ImagePayload):
            return node
    raise AssertionError("image node not found")


def _edit(node, edit_type: str, payload: dict):
    return EditOperation(
        operation_id=f"test-{edit_type}",
        type=edit_type,
        target_node_id=node.node_id,
        precondition=EditPrecondition(
            expected_semantic_digest=node_semantic_digest(node),
            expected_native_locator_digest=native_locator_digest(node),
            expected_old_value=node_semantic_text(node),
        ),
        payload=payload,
    )


def _assert_only_member_changed(before: bytes, after: bytes, changed: str):
    with ZipFile(BytesIO(before)) as a, ZipFile(BytesIO(after)) as b:
        assert a.namelist() == b.namelist()
        for name in a.namelist():
            if name == changed:
                assert a.read(name) != b.read(name)
            else:
                assert a.read(name) == b.read(name), name


def test_noop_patch_is_exact_archive_copy():
    from markitdown.twoways.formats.pptx import patch_pptx, read_pptx_ir

    source = make_pptx_bytes()
    document = read_pptx_ir(BytesIO(source))
    output = BytesIO()
    result = patch_pptx(document, BytesIO(source), output, edits=())

    assert output.getvalue() == source
    assert result.mode == "patch"
    assert result.fidelity.claimed_tier == "exact-preserve"
    assert result.metadata["touched_parts"] == ()


def test_end_to_end_text_patch_changes_only_one_slide_xml_member():
    from markitdown.twoways.formats.pptx import patch_pptx, read_pptx_ir

    source = make_pptx_bytes()
    document = read_pptx_ir(BytesIO(source))
    node = _text_node(document)
    output = BytesIO()
    result = patch_pptx(
        document,
        BytesIO(source),
        output,
        edits=(_edit(node, "replace_text", {"text": "Revenue 42%"}),),
    )
    patched = output.getvalue()

    presentation = Presentation(BytesIO(patched))
    textbox = next(shape for shape in presentation.slides[0].shapes if shape.shape_id == 3)
    assert textbox.text == "Revenue 42%"
    assert [run.text for run in textbox.text_frame.paragraphs[0].runs] == ["Revenue ", "42%"]
    assert textbox.text_frame.paragraphs[0].runs[0].font.bold is True
    assert textbox.text_frame.paragraphs[0].runs[1].font.italic is True

    _assert_only_member_changed(source, patched, "ppt/slides/slide1.xml")
    assert result.fidelity.claimed_tier == "high"
    assert result.metadata["touched_parts"] == ("ppt/slides/slide1.xml",)


def test_end_to_end_alt_text_patch_changes_only_one_slide_xml_member():
    from markitdown.twoways.formats.pptx import patch_pptx, read_pptx_ir

    source = make_pptx_bytes()
    document = read_pptx_ir(BytesIO(source))
    node = _image_node(document)
    output = BytesIO()
    patch_pptx(
        document,
        BytesIO(source),
        output,
        edits=(_edit(node, "set_alt_text", {"alt_text": "Updated icon"}),),
    )
    patched = output.getvalue()

    presentation = Presentation(BytesIO(patched))
    picture = next(shape for shape in presentation.slides[0].shapes if shape.shape_id == 4)
    assert picture._element._nvXxPr.cNvPr.get("descr") == "Updated icon"
    _assert_only_member_changed(source, patched, "ppt/slides/slide1.xml")


def test_identity_markdown_to_minimal_pptx_patch_vertical_slice():
    from markitdown.twoways import (
        MarkdownProjectionMode,
        MarkdownProjectionOptions,
        import_identity_markdown,
        project_markdown,
    )
    from markitdown.twoways.formats.pptx import patch_pptx, read_pptx_ir

    source = make_pptx_bytes()
    document = read_pptx_ir(BytesIO(source))
    projection = project_markdown(
        document,
        options=MarkdownProjectionOptions(mode=MarkdownProjectionMode.IDENTITY),
    )
    edited_markdown = projection.markdown.replace("**Revenue ***38%*", "**Revenue ***42%*", 1)
    imported = import_identity_markdown(
        edited_markdown,
        original_document=document,
        manifest=projection.manifest,
    )
    assert [edit.type for edit in imported.edits] == ["replace_text"]

    output = BytesIO()
    patch_pptx(document, BytesIO(source), output, edits=imported.edits)
    presentation = Presentation(BytesIO(output.getvalue()))
    textbox = next(shape for shape in presentation.slides[0].shapes if shape.shape_id == 3)
    assert textbox.text == "Revenue 42%"
    _assert_only_member_changed(source, output.getvalue(), "ppt/slides/slide1.xml")


def test_source_digest_mismatch_fails_before_output_is_written():
    from markitdown.twoways.formats.pptx import patch_pptx, read_pptx_ir

    source = make_pptx_bytes(metric="38%")
    wrong_source = make_pptx_bytes(metric="39%")
    document = read_pptx_ir(BytesIO(source))
    node = _text_node(document)
    output = BytesIO()

    with pytest.raises(SourcePackageMismatchError):
        patch_pptx(
            document,
            BytesIO(wrong_source),
            output,
            edits=(_edit(node, "replace_text", {"text": "Revenue 42%"}),),
        )
    assert output.getvalue() == b""


def test_end_to_end_notes_patch_changes_only_notes_slide_xml_member():
    from markitdown.twoways.formats.pptx import patch_pptx, read_pptx_ir

    source = make_pptx_bytes()
    document = read_pptx_ir(BytesIO(source))
    note = next(node for node in document.nodes.values() if node.kind == "note")
    output = BytesIO()
    patch_pptx(
        document,
        BytesIO(source),
        output,
        edits=(_edit(note, "replace_text", {"text": "Speaker note 42%"}),),
    )
    patched = output.getvalue()
    presentation = Presentation(BytesIO(patched))
    assert presentation.slides[0].notes_slide.notes_text_frame.text == "Speaker note 42%"
    _assert_only_member_changed(source, patched, "ppt/notesSlides/notesSlide1.xml")


def test_end_to_end_group_child_text_patch_preserves_group_and_sibling():
    from markitdown.twoways.formats.pptx import patch_pptx, read_pptx_ir
    from ._pptx_fixtures import make_grouped_pptx_bytes

    source = make_grouped_pptx_bytes()
    document = read_pptx_ir(BytesIO(source))
    group = next(node for node in document.nodes.values() if node.kind == "group")
    target = next(
        document.nodes[node_id]
        for node_id in group.children
        if isinstance(document.nodes[node_id].payload, TextPayload)
        and document.nodes[node_id].payload.text == "38%"
    )
    sibling = next(
        document.nodes[node_id]
        for node_id in group.children
        if node_id != target.node_id
    )

    output = BytesIO()
    patch_pptx(
        document,
        BytesIO(source),
        output,
        edits=(_edit(target, "replace_text", {"text": "42%"}),),
    )
    patched = output.getvalue()
    reread = read_pptx_ir(BytesIO(patched))

    assert reread.nodes[group.node_id].children == group.children
    assert reread.nodes[target.node_id].payload.text == "42%"
    assert node_semantic_digest(reread.nodes[sibling.node_id]) == node_semantic_digest(sibling)
    _assert_only_member_changed(source, patched, "ppt/slides/slide1.xml")


def test_verifier_rejects_unrelated_native_shape_mutation_inside_touched_slide():
    from lxml import etree

    from markitdown.twoways import RoundTripVerificationError
    from markitdown.twoways.formats.pptx import patch_pptx, read_pptx_ir
    from markitdown.twoways.formats.pptx.verify import verify_pptx_output
    from markitdown.twoways.ooxml import OOXMLPackageLimits

    source = make_pptx_bytes()
    document = read_pptx_ir(BytesIO(source))
    node = _text_node(document)
    edit = _edit(node, "replace_text", {"text": "Revenue 42%"})
    output = BytesIO()
    patch_pptx(document, BytesIO(source), output, edits=(edit,))

    with ZipFile(BytesIO(output.getvalue()), "r") as archive:
        members = {name: archive.read(name) for name in archive.namelist()}
    root = etree.fromstring(members["ppt/slides/slide1.xml"])
    picture_cnvpr = root.xpath('.//*[local-name()="pic"]//*[local-name()="cNvPr"]')[0]
    picture_cnvpr.set("name", "Tampered picture name")
    members["ppt/slides/slide1.xml"] = etree.tostring(
        root,
        xml_declaration=True,
        encoding="UTF-8",
        standalone=True,
    )
    tampered_stream = BytesIO()
    with ZipFile(tampered_stream, "w") as archive:
        for name, payload in members.items():
            archive.writestr(name, payload)

    with pytest.raises(RoundTripVerificationError) as exc_info:
        verify_pptx_output(
            document,
            source,
            tampered_stream.getvalue(),
            edits=(edit,),
            touched_parts=("ppt/slides/slide1.xml",),
            limits=OOXMLPackageLimits(),
        )
    assert exc_info.value.details["check"] == "native.unrelated_subtrees"
